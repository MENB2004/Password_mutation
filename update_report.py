import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

def update_docx_report():
    docx_path = Path("final_submission/Password_Strength_Predictable_Mutation_Final_Report.docx")
    doc = docx.Document(str(docx_path))

    # Helper function to style table headers
    def style_table(table, header_bg="1A2438", font_size=9.5):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(font_size)
                        if i == 0:
                            run.font.bold = True

    # 1. Update Section 3 (Objectives) with Verification Table
    # Find paragraph index for "4. Mathematical Foundation"
    idx_sec4 = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("4. Mathematical Foundation"):
            idx_sec4 = i
            break

    if idx_sec4 is not None:
        p_target = doc.paragraphs[idx_sec4]
        # Insert introductory text before Section 4
        p_intro = p_target.insert_paragraph_before("Project Objectives Fulfillment Matrix:")
        p_intro.runs[0].font.bold = True

        # Insert Table
        table_obj = doc.add_table(rows=8, cols=4)
        table_obj.style = "Table Grid"

        headers = ["Original Objective", "Implementation", "Evidence", "Status"]
        for col_idx, h in enumerate(headers):
            table_obj.cell(0, col_idx).text = h

        rows_data = [
            ("Measure conventional password strength", "Theoretical search-space (S = N^L) and Shannon entropy (H = L * log2(N)) model", "core.py, baseline_comparison.csv, application UI", "COMPLETED"),
            ("Detect dictionary and pattern risks", "Trie/set dictionary lookup matching embedded roots and normalized patterns", "COMMON_WORDS in core.py, test suite", "COMPLETED"),
            ("Detect predictable mutations", "Channel detectors for TitleCase, leetspeak substitution, numeric suffixes, and symbol suffixes", "detect_mutations() in core.py, dataset labels", "COMPLETED"),
            ("Develop Mutation Vulnerability Score (MVS)", "Fixed-weight MVS and Laplace-smoothed empirical MVS: MVS_emp = 100 * [sum(w_i * p_hat_i * x_i)] / [sum(w_i * x_i)]", "Formula (4.4 & 4.5), empirical_mutation_probabilities.csv", "COMPLETED"),
            ("Combine factors into realistic assessment (PSI)", "Password Security Index: PSI = 100 * E_norm * (1 - MVS/100)", "core.py, Streamlit UI, reports", "COMPLETED"),
            ("Compare against conventional & existing methods", "Empirical benchmark comparing Conventional Entropy, Dropbox zxcvbn, MVS, and PSI", "Section 8.3, results/baseline_comparison.csv, BASELINE_COMPARISON_REPORT.md", "COMPLETED"),
            ("Independent Calibration & Held-out Validation", "H_ref calibration (H_ref = 50.0) and threshold decision (tau = 69.0) on held-out split", "Phase 5 report, results/phase5_validation_metrics.csv", "COMPLETED"),
        ]

        for row_idx, rdata in enumerate(rows_data, start=1):
            for col_idx, val in enumerate(rdata):
                table_obj.cell(row_idx, col_idx).text = val

        style_table(table_obj, font_size=9)

        # Move table before Section 4
        p_target._p.addprevious(table_obj._tbl)
        doc.paragraphs[idx_sec4].insert_paragraph_before("")

    # 2. Add Section 8.3 "Comparison with Existing Password Strength Assessment"
    # Locate Section 9 ("9. Discussion") to insert Section 8.3 before it
    idx_sec9 = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("9. Discussion"):
            idx_sec9 = i
            break

    if idx_sec9 is not None:
        p_sec9 = doc.paragraphs[idx_sec9]

        p_h83 = p_sec9.insert_paragraph_before("8.3 Comparison with Existing Password Strength Assessment")
        p_h83.style = "Heading 2"

        p1 = p_sec9.insert_paragraph_before(
            "To place the proposed Mutation Vulnerability Score (MVS) and Password Security Index (PSI) in "
            "the context of established cybersecurity engineering, we conducted an empirical comparison against "
            "both conventional theoretical entropy and an established pattern-matching password estimator, Dropbox zxcvbn. "
            "The experiment evaluates whether the proposed metric captures predictable mutation structure that is not "
            "represented by conventional entropy alone, without replacing existing tools."
        )

        p2 = p_sec9.insert_paragraph_before(
            "1. Why Entropy Alone Overestimates Human-Generated Password Strength:\n"
            "Conventional password meters rely heavily on theoretical Shannon entropy under the assumption of "
            "independent uniform character selection (H = L * log2(N)). When a user constructs a password such as "
            "'Password123!', the presence of lowercase letters, uppercase letters, digits, and symbols implies an "
            "alphabet size N = 94, yielding H = 78.66 bits and an apparent search space of 4.75 x 10^23 combinations. "
            "In practice, an attacker using dictionary and rule-based mutation engines (e.g. Hashcat, John the Ripper) "
            "does not search all 94^12 arbitrary combinations. Instead, the search is constrained to [Common Root] + [Number Suffix] + [Symbol], "
            "reducing the effective guessing space to thousands rather than sextillions. Conventional entropy fails to reflect this reduction."
        )

        p3 = p_sec9.insert_paragraph_before(
            "2. How Existing Password-Strength Estimators Approach the Problem:\n"
            "Dropbox zxcvbn addresses this limitation by using pattern matching against extensive built-in dictionaries "
            "(common passwords, names, surnames, English words, popular culture terms) and detecting spatial keyboard walks "
            "and l33t substitutions. It models password guessing entropy using a dynamic programming path-search to find "
            "the shortest sequence of patterns that reproduces the password, reporting an integer score from 0 to 4 along with "
            "estimated offline cracking times."
        )

        p4 = p_sec9.insert_paragraph_before(
            "3. What Information MVS Adds:\n"
            "While zxcvbn produces a holistic score based on minimum path guesses, the proposed MVS provides an explicitly "
            "auditable, feature-by-feature decomposition: MVS_emp = 100 * [sum(w_i * p_hat_i * x_i)] / [sum(w_i * x_i)]. "
            "This separates the evaluation into transparent, interpretable channels: dictionary-base dependence, first-letter "
            "or TitleCase capitalization, character substitutions, numeric suffixes, and symbol additions. Each mutation channel "
            "is grounded in Laplace-smoothed empirical frequencies (p_hat = (k+1)/(n+2)) derived from observed distributions."
        )

        p5 = p_sec9.insert_paragraph_before(
            "4. How PSI Combines Entropy and Predictable Mutation:\n"
            "The Password Security Index (PSI) combines theoretical character-space capacity with the mutation penalty: "
            "PSI = 100 * min(H / H_ref, 1.0) * (1 - MVS / 100). If a password contains genuine, non-dictionary random diversity "
            "(such as 'vQ7mK2xR9zP4'), MVS evaluates to 0.0, and the password retains its full normalized entropy. If predictable "
            "mutations are detected, the normalized entropy is proportionally discounted."
        )

        p6 = p_sec9.insert_paragraph_before("5. Experimental Benchmark Results:")
        p6.runs[0].font.bold = True

        # Insert Benchmark Table
        table_bench = doc.add_table(rows=6, cols=7)
        table_bench.style = "Table Grid"
        bench_headers = ["Assessment Method", "Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC", "PR-AUC"]
        for c_idx, bh in enumerate(bench_headers):
            table_bench.cell(0, c_idx).text = bh

        bench_data = [
            ("Proposed PSI (tau=69.0, H_ref=50)", "0.9788", "1.0000", "0.9333", "0.9655", "1.0000", "1.0000"),
            ("Proposed MVS (> 0)", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000"),
            ("Proposed Empirical MVS", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000", "1.0000"),
            ("Dropbox zxcvbn (Score <= 2)", "0.9242", "0.8077", "1.0000", "0.8936", "0.9937", "0.9780"),
            ("Conventional Entropy (< Median)", "0.8000", "0.6211", "0.9524", "0.7519", "0.9365", "0.8913"),
        ]
        for r_idx, rvals in enumerate(bench_data, start=1):
            for c_idx, val in enumerate(rvals):
                table_bench.cell(r_idx, c_idx).text = val

        style_table(table_bench, font_size=9)
        p_sec9._p.addprevious(table_bench._tbl)

        p7 = p_sec9.insert_paragraph_before(
            "6. Case Study Analysis: Three Canonical Archetypes:\n"
            "To illustrate how each model behaves across different structural patterns, Table 8.4 compares three test passwords:\n"
            "• Password A ('Password123!'): Demonstrates the classic predictable suffix mutation. Conventional entropy awards "
            "100/100 (78.66 bits), yet MVS flags 80.0/100 and discounts PSI to 19.7/100. Dropbox zxcvbn concurs, rating it 1/4 (Weak, ~4s crack time).\n"
            "• Password B ('vQ7mK2xR9zP4'): A synthetic random-like control string. Having no dictionary base, MVS correctly evaluates to 0.0/100, "
            "preserving full entropy (PSI: 89.3/100, Calibrated PSI: 100/100). Zxcvbn rates it 4/4 (Strong, ~3 years crack time).\n"
            "• Password C ('P@ssword2026!'): A multi-layered mutation combining TitleCase, leetspeak substitution (@ for a), a 4-digit year suffix (2026), "
            "and a symbol (!). Conventional entropy awards 100/100 (85.21 bits). However, all 4 mutation channels are detected, resulting in MVS: 100.0/100 "
            "and PSI: 0.0/100. Zxcvbn rates it 2/4 (Fair, ~17 minutes crack time)."
        )

        p8 = p_sec9.insert_paragraph_before(
            "7. Limitations of Synthetic Data Comparison:\n"
            "The observed classification separation (ROC-AUC 1.000 for PSI and 0.994 for zxcvbn) describes the controlled synthetic "
            "research dataset. Because the dataset was constructed to test these specific mutation families, the high metric performance "
            "reflects internal consistency rather than guaranteed real-world cracking accuracy. In real-world environments, passwords exhibit "
            "linguistic, multilingual, and contextual mutations beyond the current small dictionary."
        )

    # 3. Update Section 10 (Limitations)
    idx_sec10 = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("10. Limitations"):
            idx_sec10 = i
            break

    if idx_sec10 is not None:
        p_sec10 = doc.paragraphs[idx_sec10]
        p_sec10_body = p_sec10.insert_paragraph_before(
            "The findings and metrics presented in this project must be interpreted in light of the following scientific constraints:\n"
            "1. Experimental Academic Metrics: MVS and PSI are project-defined research metrics and do not represent NIST, ISO, or industry standards.\n"
            "2. Controlled Synthetic Data: All statistical validation was conducted on synthetic datasets constructed to test specific hypothesis families; "
            "results cannot be directly extrapolated to general human populations without empirical validation on licensed real breach corpora.\n"
            "3. Theoretical Entropy Limitations: Search space (S = N^L) and Shannon entropy (H = L * log2(N)) represent idealized upper bounds and do not measure true attacker resistance.\n"
            "4. Dictionary and Rule Scope: The mutation detector currently operates on a compact educational dictionary (24 roots) and four specific mutation channels. "
            "Attacks in the wild exploit slang, multi-word phrases, personal data, and complex token sequences.\n"
            "5. Cracking Outcomes: PSI does not calculate exact seconds-to-crack. Offline crack-times estimated by existing tools (e.g. zxcvbn) are hardware-dependent benchmarks.\n"
            "6. Complementary Role: The proposed system is intended to provide interpretable structural transparency into password construction habits, "
            "complementing rather than replacing established pattern estimators."
        )

    doc.save(str(docx_path))
    print(f"Successfully updated DOCX report: {docx_path}")

if __name__ == "__main__":
    update_docx_report()
