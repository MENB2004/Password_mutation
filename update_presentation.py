import sys
from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def update_presentation():
    pptx_path = Path("final_submission/Password_Strength_Predictable_Mutation_Presentation.pptx")
    prs = pptx.Presentation(str(pptx_path))
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_title = RGBColor(26, 36, 56)      # Deep Navy
    c_body = RGBColor(50, 50, 50)       # Charcoal
    c_highlight = RGBColor(0, 102, 204) # Blue

    def add_standard_slide(title_text, bullet_items):
        slide = prs.slides.add_slide(blank_layout)

        # Title Box
        t_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.8))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = c_title

        # Body Box
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.0))
        bf = b_box.text_frame
        bf.word_wrap = True

        for i, item in enumerate(bullet_items):
            p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            p.space_after = Pt(10)

            if isinstance(item, tuple):
                header, text = item
                run1 = p.add_run()
                run1.text = header + ": "
                run1.font.bold = True
                run1.font.size = Pt(17)
                run1.font.color.rgb = c_title

                run2 = p.add_run()
                run2.text = text
                run2.font.size = Pt(16)
                run2.font.color.rgb = c_body
            else:
                p.text = "• " + item
                p.font.size = Pt(17)
                p.font.color.rgb = c_body

        return slide

    # SLIDE 1: Comparison with Existing Password Strength Assessment
    s1_bullets = [
        ("Conventional Entropy Limitation", "Assumes characters are chosen independently and uniformly at random (H = L * log2(N)). Severely overestimates human-constructed passwords."),
        ("Existing Estimator Approach (Dropbox zxcvbn)", "Employs dictionary matching, spatial keyboard walks, and l33t substitutions. Estimates cracking guesses and realistic offline crack times."),
        ("Proposed Mutation Vulnerability (MVS)", "Provides an explicit, transparent decomposition: MVS = 100 * [sum(w_i * p_hat_i * x_i)] / [sum(w_i * x_i)]. Quantifies specific predictable human construction habits."),
        ("Combined Security Index (PSI)", "Integrates Shannon entropy with mutation penalties: PSI = 100 * E_norm * (1 - MVS/100). Preserves character diversity while penalizing predictable roots."),
        ("Scientific Perspective", "MVS/PSI and zxcvbn offer complementary insights: zxcvbn focuses on guessability, while MVS provides an auditable, component-by-component vulnerability score.")
    ]
    add_standard_slide("Comparison with Existing Password Strength Assessment", s1_bullets)

    # SLIDE 2: Experimental Baseline Comparison
    s2_bullets = [
        ("Controlled Synthetic Benchmark (n = 330)", "105 predictable mutation instances across 7 families + 225 synthetic random-like controls."),
        ("Proposed PSI Model", "Accuracy: 97.88% | Precision: 1.0000 | Recall: 0.9333 | ROC-AUC: 1.0000 | PR-AUC: 1.0000 (Calibrated at H_ref = 50, tau = 69.0)."),
        ("Proposed MVS Model", "Accuracy: 100.0% | Precision: 1.0000 | Recall: 1.0000 | ROC-AUC: 1.0000 | PR-AUC: 1.0000 (MVS > 0 separates all families)."),
        ("Dropbox zxcvbn (Score <= 2 as Vulnerable)", "Accuracy: 92.42% | Precision: 0.8077 | Recall: 1.0000 | ROC-AUC: 0.9937 | PR-AUC: 0.9780."),
        ("Conventional Entropy-Only (< Median)", "Accuracy: 80.00% | Precision: 0.6211 | Recall: 0.9524 | ROC-AUC: 0.9365 | PR-AUC: 0.8913."),
        ("Key Finding", "Conventional entropy shows high false-positive security for predictable passwords. Both PSI and zxcvbn successfully detect dictionary-dependent mutations.")
    ]
    add_standard_slide("Experimental Baseline Comparison", s2_bullets)

    # SLIDE 3: Example: Entropy vs Predictable Mutation
    s3_bullets = [
        ("Password A — 'Password123!' (Predictable Suffix)", "Length: 12 | Pool: 94 | Conventional Meter: 100/100 | Entropy: 78.66 bits | MVS: 80.0/100 | Proposed PSI: 19.7/100 (High Risk) | zxcvbn: 1/4 (Weak, ~4s crack time)."),
        ("Password B — 'vQ7mK2xR9zP4' (Random-Like Control)", "Length: 12 | Pool: 62 | Conventional Meter: 81.2/100 | Entropy: 71.45 bits | MVS: 0.0/100 | Proposed PSI: 89.3/100 (Higher Security) | zxcvbn: 4/4 (Strong, ~3y crack time)."),
        ("Password C — 'P@ssword2026!' (Complex Multi-Layer Mutation)", "Length: 13 | Pool: 94 | Conventional Meter: 100/100 | Entropy: 85.21 bits | MVS: 100.0/100 | Proposed PSI: 0.0/100 (High Risk) | zxcvbn: 2/4 (Fair, ~17m crack time)."),
        ("Core Takeaway", "Higher theoretical entropy does NOT mean higher security when predictable mutations are present. The proposed model correctly penalizes dictionary-derived passwords while protecting random controls.")
    ]
    add_standard_slide("Example: Entropy vs Predictable Mutation", s3_bullets)

    # SLIDE 4: Final System and Contributions
    s4_bullets = [
        ("Core Analysis Engine (`password_analyzer`)", "Implements theoretical search space, entropy, fixed MVS, empirical probability MVS, PSI, and calibrated classification."),
        ("Baseline Comparator & Public Dataset Pipeline", "Local integration with Dropbox zxcvbn; privacy-preserving CSV importer computing aggregate telemetry without exposing raw credentials."),
        ("Reproducible Experimental Suite (`experiments.py`)", "Automated pipeline generating baseline comparison CSV, formal Markdown reports, and 8 publication-quality figures."),
        ("Interactive Streamlit Application (`app.py`)", "Comprehensive UI with Model Comparison deep dive, 3-password benchmark, empirical surprisal tables, visualization browser, and math formulations."),
        ("Verification & Academic Deliverables", "8/8 unit tests passing; updated DOCX academic report; complete mathematical documentation and project summary.")
    ]
    add_standard_slide("Final System and Contributions", s4_bullets)

    # Save presentation
    prs.save(str(pptx_path))
    print(f"Successfully updated presentation: {pptx_path} (Total slides: {len(prs.slides)})")

if __name__ == "__main__":
    update_presentation()
